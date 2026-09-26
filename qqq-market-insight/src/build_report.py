"""Build the 10-page QQQ course report from saved five-year pipeline evidence."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "outputs" / "tables" / "5_year"
FIGURES = ROOT / "outputs" / "figures" / "5_year"
OUTPUT = ROOT / "deliverables" / "report" / "QQQ_Technical_Indicator_Report.docx"


def rows(name: str) -> list[dict[str, str]]:
    with (TABLES / name).open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


quality = json.loads((TABLES / "data_quality.json").read_text(encoding="utf-8"))
metrics = rows("model_metrics.csv")
metric = {record["model"]: record for record in metrics}
importance = rows("permutation_importance.csv")
families = rows("group_importance.csv")
cv = {record["metric"]: record for record in rows("timeseries_cv_summary.csv")}


def pct(value: str | float, digits: int = 2) -> str:
    return f"{float(value) * 100:.{digits}f}%"


def style_font(style, size: float, bold: bool = False) -> None:
    style.font.name = "Sarabun"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), "Sarabun")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.70)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.header_distance = Inches(0.30)
section.footer_distance = Inches(0.35)

styles = doc.styles
style_font(styles["Normal"], 9.4)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.18
style_font(styles["Title"], 20, True)
styles["Title"].paragraph_format.space_after = Pt(13)
title_properties = styles["Title"].element.get_or_add_pPr()
for border in title_properties.findall(qn("w:pBdr")):
    title_properties.remove(border)
style_font(styles["Heading 1"], 15, True)
styles["Heading 1"].paragraph_format.space_before = Pt(0)
styles["Heading 1"].paragraph_format.space_after = Pt(9)
style_font(styles["Heading 2"], 10.8, True)
styles["Heading 2"].paragraph_format.space_before = Pt(9)
styles["Heading 2"].paragraph_format.space_after = Pt(4)

header = section.header.paragraphs[0]
header.text = "QQQ Technical Indicator Feature Importance"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.style = "Normal"
header.runs[0].font.size = Pt(7.5)
header.runs[0].font.color.rgb = RGBColor(100, 100, 100)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("D6900557  |  Big Data Analytics  |  ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(7.5)


def heading(text: str, level: int = 1) -> None:
    doc.add_heading(text, level)


def paragraph(text: str, *, bold_lead: str | None = None) -> None:
    item = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        item.add_run(bold_lead).bold = True
        item.add_run(text[len(bold_lead):])
    else:
        item.add_run(text)


def simple_table(headers: list[str], data: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for j, (value, width) in enumerate(zip(headers, widths)):
        table.columns[j].width = Inches(width)
        table.rows[0].cells[j].width = Inches(width)
        table.rows[0].cells[j].text = value
    for i, values in enumerate(data):
        cells = table.add_row().cells
        for j, (value, width) in enumerate(zip(values, widths)):
            cells[j].width = Inches(width)
            cells[j].text = str(value)
        if i % 2 == 1:
            for cell in cells:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F0F4F7")
                cell._tc.get_or_add_tcPr().append(shd)
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{edge}")
                item.set(qn("w:val"), "single")
                item.set(qn("w:sz"), "4")
                item.set(qn("w:color"), "D9D9D9")
                borders.append(item)
            tcpr.append(borders)
            mar = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{edge}")
                item.set(qn("w:w"), "80")
                item.set(qn("w:type"), "dxa")
                mar.append(item)
            tcpr.append(mar)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.10
                for run in p.runs:
                    run.font.name = "Sarabun"
                    run.font.size = Pt(8.2)
                    if i == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "18354B")
                tcpr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def figure(filename: str, caption: str, width: float, max_height: float | None = None) -> None:
    image_path = FIGURES / filename
    from PIL import Image

    with Image.open(image_path) as image:
        ratio = image.height / image.width
    if max_height and width * ratio > max_height:
        width = max_height / ratio
    item = doc.add_paragraph()
    item.alignment = WD_ALIGN_PARAGRAPH.CENTER
    item.paragraph_format.space_after = Pt(2)
    item.add_run().add_picture(str(image_path), width=Inches(width))
    caption_p = doc.add_paragraph(caption)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(6)
    for run in caption_p.runs:
        run.font.size = Pt(7.8)
        run.font.color.rgb = RGBColor(70, 70, 70)


def next_page() -> None:
    doc.add_page_break()


# Page 1: cover and executive summary.
doc.add_paragraph("QQQ Technical Indicator Feature Importance", "Title")
paragraph("Big Data Analytics Project Report")
paragraph("Author: Tien Hanprab  |  Student ID: D6900557")
paragraph("Source: Nasdaq Historical Quotes for QQQ  |  Data: 4 January 2021–17 September 2026")
paragraph("Prepared: 26 September 2026")
heading("Executive summary", 1)
paragraph(
    "This project examines price, volume, and volatility patterns in QQQ and tests whether 21 "
    "technical indicators can classify the following trading day's closing-price direction. "
    "Nasdaq supplied 1,433 daily observations; 1,412 remained after feature preparation [1]."
)
paragraph(
    "On a chronological test period of 283 days, the majority-class baseline achieved 54.77% "
    "accuracy, versus 51.59% for Random Forest and 53.00% for HistGradientBoosting. "
    "Neither fitted model outperformed the baseline on test accuracy."
)
paragraph(
    "HistGradientBoosting relied most on atr_14_pct in test-set permutation importance: "
    "permuting it reduced balanced accuracy by 0.0315 ± 0.0085. This measures model dependence "
    "in this sample, not causation or a profitable trading rule."
)
heading("Scope and evidence", 2)
paragraph("This report uses the executed notebooks/pipeline_qqq_5_year.ipynb and saved five-year pipeline tables and figures. It is for education, not investment advice.")
next_page()

# Page 2: introduction.
heading("1 Introduction and objectives")
paragraph(
    "QQQ is an exchange-traded fund associated with the Nasdaq-100. Its daily prices reflect "
    "longer-term trends alongside short-term volatility. This study explores observable patterns "
    "in daily prices and trading volume, then tests how useful indicators available at the close "
    "of day t are for classifying the next trading day."
)
heading("Objectives", 2)
paragraph("First, describe price, return, volume, and volatility patterns over the configured period.")
paragraph("Second, assess indicator importance for next-day classification using test-set permutation importance.")
paragraph("Third, compare Random Forest and other models with a majority-class baseline on chronologically held-out observations.")
heading("Research questions", 2)
paragraph("RQ1  What are the major price, return, volume, and volatility patterns of QQQ during the configured analysis period?")
paragraph("RQ2  Which technical indicators are the most important features for classifying QQQ's next-day price direction?")
paragraph("RQ3  Does the Random Forest classifier perform better than a simple baseline on unseen chronological data?")
heading("Target definition", 2)
paragraph("Define target_up(t) = 1 if Close(t+1) > Close(t), and 0 otherwise. Here t+1 means the next trading day, not the next calendar day. The next_return column creates the target and supports exploratory comparisons; it is never supplied to a model as a predictor.")
paragraph("The study evaluates classification, not trading costs, dividends, or realized strategy returns.")
next_page()

# Page 3: data.
heading("2 Data description and source")
paragraph(
    "Daily QQQ observations came from Nasdaq's public historical quote service via "
    "https://api.nasdaq.com/api/quote/QQQ/historical with assetclass=etf, "
    "fromdate=2021-01-01, todate=2026-09-17, and limit=5000. The Nasdaq QQQ Historical "
    "Data page is the human-readable source reference [1]. Raw JSON and retrieval metadata "
    "are stored separately; the downloaded response is preserved unchanged."
)
simple_table(
    ["Field", "Parsed type", "Meaning and units"],
    [
        ["Date", "datetime", "Trading date"],
        ["Open", "float", "Opening price, USD per share"],
        ["High", "float", "Daily high, USD per share"],
        ["Low", "float", "Daily low, USD per share"],
        ["Close", "float", "Closing price, USD per share"],
        ["Volume", "integer", "Shares traded"],
    ],
    [1.0, 1.25, 4.75],
)
paragraph(
    "The response contains 1,433 trading days from 2021-01-04 through 2026-09-17, "
    "exceeding the assignment's 500-row minimum [4]. Metadata records retrieval at "
    "2026-09-19T22:15:36Z, the request URL, and SHA-256 "
    "5af8f404421142e426ca6857137b9355c4bee85fd8067ae6923973e3dab7001a"
)
paragraph("Price-definition limitation: the report uses Close as supplied by Nasdaq. It does not call this adjusted close or total return because dividend adjustment is not verified.")
next_page()

# Page 4: cleaning.
heading("3 Data-cleaning process")
paragraph(
    "The pipeline checks the JSON schema, parses MM/DD/YYYY dates, and removes currency "
    "symbols and thousands separators before converting OHLCV fields to numbers. It then sorts "
    "oldest to newest and audits duplicate dates, missing or non-finite values, nonpositive "
    "prices, negative volume, and daily High/Low consistency."
)
simple_table(
    ["Stage", "Rows", "Explanation"],
    [
        ["Raw response", "1,433", "Saved Nasdaq observations"],
        ["Invalid OHLCV", "0", "No missing or impossible observations"],
        ["Duplicate dates", "0", "None removed"],
        ["Clean OHLCV", "1,433", "Chronologically sorted"],
        ["Indicator warm-up", "20", "Insufficient 20-day history"],
        ["Unknown next target", "1", "No following close on final day"],
        ["Model-ready", "1,412", "21 features plus next_return and target_up"],
    ],
    [1.9, 1.0, 4.1],
)
paragraph(
    "The row audit reconciles exactly: 1,433 − 20 − 1 = 1,412. No feature values are missing "
    "after warm-up removal. If an invalid OHLCV row appears in a later download, the pipeline "
    "records it and stops modeling rather than silently allowing the next-day label to skip a date."
)
paragraph("The model-ready data span 2021-02-02 through 2026-09-16 and are saved at data/processed/5_year/qqq_features.csv.")
next_page()

# Page 5: EDA I.
heading("4 Exploratory data analysis")
paragraph("The six EDA chart types on pages 5–7 use only the training period: 1,128 trading days from 2021-02-02 through 2025-07-30. Holdout outcomes were not used for this exploration.")
figure("01_price_trend.png", "Figure 1. Closing-price line and 50-day moving average in the training period.", 6.8, 2.85)
paragraph("The training-period close rose from USD 327.68 to 568.02 (+73.35%), but the path was uneven. The observed minimum was USD 260.10 on 2022-12-28, and the maximum was USD 568.14 on 2025-07-28. The long-run rise therefore coexisted with substantial short-run variation.")
figure("02_return_distribution.png", "Figure 2. Histogram of daily training-period returns.", 5.8, 2.65)
paragraph("Mean daily return was 0.0609% with a standard deviation of 1.4621%; observed returns ranged from −6.2109% to +12.0031%. The tails and extreme days show why the mean alone cannot summarize daily risk.")
next_page()

# Page 6: EDA II.
heading("4 Exploratory data analysis (continued)")
figure("03_monthly_boxplot.png", "Figure 3. Box plot of daily returns by calendar month in the training period.", 6.2, 2.85)
paragraph("Monthly medians and spreads differ, with outliers in several months. This does not establish a tradable seasonal pattern: stability across years and out-of-sample performance were not tested.")
figure("04_rsi_scatter.png", "Figure 4. Scatter plot of 14-day RSI against next-day return.", 5.8, 2.60)
paragraph("Returns fall above and below zero across RSI levels. Pearson correlation between RSI and next_return is −0.0312, showing little linear association from this indicator alone.")
next_page()

# Page 7: EDA III.
heading("4 Exploratory data analysis (continued)")
figure("05_correlation_heatmap.png", "Figure 5. Pearson correlation heatmap for technical indicators.", 5.3, 4.45)
paragraph("Some indicators are highly correlated: ema_ratio_12 versus sma_ratio_10 has r = 0.9694, and macd versus macd_signal has r = 0.9506. Redundancy can depress individual permutation importance when a correlated substitute remains available [3].")
figure("06_class_balance.png", "Figure 6. Bar chart of the training-period target classes.", 4.4, 2.25)
paragraph("Training data contain 614 Up days (54.43%) and 514 Down-or-flat days (45.57%). This imbalance makes balanced accuracy a useful complement to raw accuracy.")
next_page()

# Page 8: methods.
heading("5 Analytical methods")
paragraph("Feature engineering derives 21 indicators from OHLCV using information available no later than the close of day t. They comprise return history (3), trend (7), momentum (3), volatility (5), and volume (3) features. Examples include return_1d, SMA/EMA ratios, RSI 14, ATR 14, and 20-day volume ratio.")
paragraph("Four fixed-configuration models were compared: a majority-class baseline; Logistic Regression with median imputation and scaling within its pipeline; Random Forest with 700 trees, max_depth 5, and min_samples_leaf 12; and HistGradientBoosting with 200 iterations and learning_rate 0.03. Output tables label the latter GradientBoosting.")
simple_table(
    ["Partition", "Rows", "Date range"],
    [
        ["Train", "1,128", "2021-02-02 to 2025-07-30"],
        ["Boundary gap", "1", "2025-07-31"],
        ["Test", "283", "2025-08-01 to 2026-09-16"],
    ],
    [1.7, 1.0, 4.3],
)
paragraph("Train/test partitioning follows time order and leaves one boundary row unused so the final training label is realized before test begins. Training validation uses expanding TimeSeriesSplit with five folds and a one-row gap [2]. There is no random shuffle.")
paragraph("The held-out test evaluation reports accuracy, balanced accuracy, precision, recall, F1, and ROC AUC. GradientBoosting permutation importance shuffles each test feature 30 times and measures the resulting reduction in balanced accuracy. A larger positive value means greater model dependence in this sample [3].")
next_page()

# Page 9: results.
heading("6 Results and discussion")
simple_table(
    ["Model", "Accuracy", "Balanced", "F1", "ROC AUC"],
    [
        [name, pct(metric[name]["accuracy"]), pct(metric[name]["balanced_accuracy"]), pct(metric[name]["f1"]), f'{float(metric[name]["roc_auc"]):.4f}']
        for name in ("Majority baseline", "Logistic regression", "Random forest", "GradientBoosting")
    ],
    [2.2, 1.1, 1.15, 1.05, 1.5],
)
paragraph("RQ3: Random Forest accuracy was 51.59%, 3.18 percentage points below the 54.77% baseline. Its balanced accuracy was 50.98% and ROC AUC 0.5034, providing no clear evidence of superior discrimination.")
paragraph("GradientBoosting had the best balanced accuracy among fitted models at 52.74%, but its 53.00% accuracy remained 1.77 points below baseline. ROC AUC was only 0.5269; this modest ranking ability is not a reliable forecasting system.")
paragraph(f"Within-training GradientBoosting cross-validation yielded mean balanced accuracy {pct(cv['balanced_accuracy']['mean'])} ± {pct(cv['balanced_accuracy']['standard_deviation'])} across folds and mean ROC AUC {float(cv['roc_auc']['mean']):.4f} ± {float(cv['roc_auc']['standard_deviation']):.4f}. These near-chance averages qualify any interpretation of the single holdout result.")
figure("07_feature_importance.png", "Figure 7. Test-set GradientBoosting permutation importance; error bars show one shuffle SD.", 6.5, 2.45)
paragraph("RQ2: atr_14_pct ranked first at 0.0315 ± 0.0085, followed by rsi_14 at 0.0127 ± 0.0072 and roc_10 at 0.0121 ± 0.0109. Summed importance was highest for the volatility family (0.0577), although this sum does not jointly permute the family and correlated variables may share importance.")
next_page()

# Page 10: conclusion, recommendations, references.
heading("7 Conclusion and recommendations")
paragraph("RQ1: Training-period prices trended upward, but daily returns were volatile and included outliers. Mean volume was about 49.89 million shares per day. EDA does not suggest that price trend or RSI alone provides an easy next-day directional signal.")
paragraph("RQ2: For GradientBoosting, 14-day ATR and the volatility family had the greatest test-set predictive association. This finding depends on the model, period, scoring metric, and correlations among indicators.")
paragraph("RQ3: Random Forest did not beat the majority baseline on accuracy in 283 test days. GradientBoosting's balanced accuracy was only slightly above 50%, while cross-validation averages were near chance. These findings do not establish a profitable trading strategy.")
heading("Recommendations", 2)
paragraph("Future work should pre-specify walk-forward evaluation across multiple periods and test whether importance rankings persist after redundant indicators are reduced. A genuine investment study would also need adjusted prices or dividends, transaction costs, and net-of-cost evaluation.")
heading("Limitations", 2)
paragraph("The data comprise daily OHLCV for one asset and a Close series not verified as dividend-adjusted. Because this test period has already been inspected, comparisons are exploratory rather than confirmation on untouched data. Feature importance is not causal evidence.")
heading("8 References")
paragraph("[1] Nasdaq. Invesco QQQ Trust Series 1 Historical Data. https://www.nasdaq.com/market-activity/etf/qqq/historical ; API: https://api.nasdaq.com/api/quote/QQQ/historical (accessed 26 September 2026).")
paragraph("[2] scikit-learn developers. TimeSeriesSplit documentation. https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.TimeSeriesSplit.html (accessed 26 September 2026).")
paragraph("[3] scikit-learn developers. Permutation feature importance documentation. https://scikit-learn.org/stable/modules/permutation_importance.html (accessed 26 September 2026).")
paragraph("[4] Big Data Analytics Project Assignment. Course requirements in docs/Big_Data_Analytics_Project_Assignment.pdf.")
paragraph("Computational evidence: notebooks/pipeline_qqq_5_year.ipynb, outputs/tables/5_year/, and outputs/figures/5_year/.")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
